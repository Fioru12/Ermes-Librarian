from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest

from core.document_parser import (
    DocumentParseError,
    chunk_source_units,
    extract_source_units,
    extract_text,
    split_into_chunks,
)


def test_extracts_text_and_markdown():
    text, units = extract_text("procedura.md", b"# Procedura\n\nContattare HR")

    assert text == "Contattare HR"
    assert units == 1


def test_extracts_shared_strings_from_xlsx():
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types />")
        archive.writestr(
            "xl/sharedStrings.xml",
            '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><t>Scadenza</t></si><si><t>31 dicembre</t></si></sst>',
        )
        archive.writestr(
            "xl/workbook.xml",
            '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Scadenze" sheetId="1" r:id="rId1" /></sheets></workbook>',
        )
        archive.writestr(
            "xl/_rels/workbook.xml.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Target="worksheets/sheet1.xml" /></Relationships>',
        )
        archive.writestr(
            "xl/worksheets/sheet1.xml",
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>1</v></c></row></sheetData></worksheet>',
        )

    text, units = extract_text("scadenze.xlsx", buffer.getvalue())

    assert text == "A1: Scadenza | B1: 31 dicembre"
    assert units == 1


def test_rejects_unsupported_file_type():
    with pytest.raises(DocumentParseError):
        extract_text("immagine.png", b"not an image")


def _write_minimal_pptx(archive: ZipFile) -> None:
    """Build a deterministic two-slide PPTX skeleton for tests."""
    a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    archive.writestr("[Content_Types].xml", "<Types />")
    archive.writestr(
        "ppt/presentation.xml", '<presentation xmlns="http://schemas.openxmlformats.org/presentationml/2006/main" />'
    )
    archive.writestr(
        "ppt/slides/slide1.xml",
        f'<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="{a}">'
        f"<p:txBody><a:p><a:r><a:t>Policy Ferie</a:t></a:r></a:p>"
        f"<a:p><a:r><a:t>Massimo </a:t></a:r><a:r><a:t>15 giorni</a:t></a:r></a:p></p:txBody></p:sld>",
    )
    archive.writestr(
        "ppt/slides/slide2.xml",
        f'<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="{a}">'
        f"<p:txBody><a:p><a:r><a:t>Note spese</a:t></a:r></a:p></p:txBody></p:sld>",
    )


def test_extracts_pptx_slides_with_locators():
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        _write_minimal_pptx(archive)

    units = extract_source_units("presentazione.pptx", buffer.getvalue())

    assert len(units) == 2
    assert units[0].locator == "Slide 1"
    assert units[0].text == "Policy Ferie\nMassimo 15 giorni"  # runs della stessa riga uniti
    assert units[1].locator == "Slide 2"
    assert units[1].text == "Note spese"


def test_rejects_a_zip_disguised_as_pptx():
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("note.txt", "non è una presentazione")

    with pytest.raises(DocumentParseError):
        extract_text("falsa.pptx", buffer.getvalue())


def test_pptx_with_entity_declarations_is_rejected():
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        _write_minimal_pptx(archive)
        archive.writestr(
            "ppt/slides/slide9.xml",
            '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
            '<!DOCTYPE p:sld [<!ENTITY xxe "boom">]></p:sld>',
        )

    with pytest.raises(DocumentParseError):
        extract_text("maligna.pptx", buffer.getvalue())


def test_rejects_a_zip_disguised_as_xlsx():
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("note.txt", "non è un foglio Excel")

    with pytest.raises(DocumentParseError):
        extract_text("falso.xlsx", buffer.getvalue())


def test_chunks_keep_paragraphs_and_bound_size():
    chunks = split_into_chunks("Uno.\n\n" + ("Due " * 400), max_chars=100, overlap_chars=20)

    assert len(chunks) > 2
    assert all(len(chunk) <= 100 for chunk in chunks)


def test_markdown_sections_are_preserved_as_chunk_citations():
    units = extract_source_units(
        "policy.md", b"# Ferie\nRichiedere ferie via portale.\n\n# Spese\nConservare la ricevuta."
    )
    chunks = chunk_source_units(units)

    assert chunks[0][1] == "Sezione: Ferie"
    assert chunks[1][1] == "Sezione: Spese"


def test_public_demo_corpus_is_parseable_and_has_section_locators():
    corpus = Path(__file__).parents[1] / "examples" / "demo-corpus"
    documents = sorted(corpus.glob("*.md"))

    assert {document.name for document in documents} >= {
        "employee-handbook.md",
        "expense-policy.md",
        "it-access-policy.md",
    }
    for document in documents:
        units = extract_source_units(document.name, document.read_bytes())
        chunks = chunk_source_units(units)

        assert chunks
        assert all(locator.startswith("Sezione:") for _, locator in chunks)


def test_xlsx_with_entity_declarations_is_rejected():
    """An uploaded workbook must not be able to expand XML entities.

    `xml.etree.ElementTree` expands internal entities, so a few kilobytes of
    nested definitions expand to gigabytes in memory (the "billion laughs"
    denial of service). Measured before the guard existed: 1.157 bytes of
    input produced 10.001 characters with only four nesting levels.
    """
    bomb = (
        '<?xml version="1.0"?>\n'
        "<!DOCTYPE sst [\n"
        '  <!ENTITY a "AAAAAAAAAA">\n'
        '  <!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;">\n'
        '  <!ENTITY c "&b;&b;&b;&b;&b;&b;&b;&b;&b;&b;">\n'
        "]>\n"
        '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="1">'
        "<si><t>&c;</t></si></sst>"
    )
    workbook = (
        '<?xml version="1.0"?>'
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<sheets><sheet name="Foglio1" sheetId="1"/></sheets></workbook>'
    )
    sheet = (
        '<?xml version="1.0"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<sheetData><row r="1"><c r="A1" t="s"><v>0</v></c></row></sheetData></worksheet>'
    )

    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types />")
        archive.writestr("xl/workbook.xml", workbook)
        archive.writestr("xl/sharedStrings.xml", bomb)
        archive.writestr("xl/worksheets/sheet1.xml", sheet)

    with pytest.raises(DocumentParseError):
        extract_source_units("bomba.xlsx", buffer.getvalue())


def test_extracts_csv_units_with_header_mapping():
    csv_bytes = b"Prodotto,Reparto,Prezzo\nLaptop,IT,999\nMouse,IT,25"
    text, units = extract_text("inventario.csv", csv_bytes)
    assert units == 2
    assert "Prodotto: Laptop | Reparto: IT | Prezzo: 999" in text


def test_extracts_rtf_units():
    rtf_bytes = (
        b"{\\rtf1\\ansi\\deff0 {\\fonttbl {\\f0 Courier;}}\\f0\\fs24 Procedura di emergenza: evacuare l'edificio.\\par}"
    )
    text, units = extract_text("emergenza.rtf", rtf_bytes)
    assert units == 1
    assert "evacuare l'edificio" in text


def test_chunk_size_configuration_is_actually_honoured(monkeypatch):
    """ERMES_CHUNK_SIZE ed ERMES_CHUNK_OVERLAP erano dichiarati nel config e
    letti da nessuna riga di codice: chi li impostava non cambiava niente."""
    import config
    from core.document_parser import split_into_chunks

    testo = "frase di prova. " * 400

    monkeypatch.setattr(config, "cfg", config.cfg.replace(CHUNK_SIZE=300, CHUNK_OVERLAP=50))
    stretti = split_into_chunks(testo)

    monkeypatch.setattr(config, "cfg", config.cfg.replace(CHUNK_SIZE=1500, CHUNK_OVERLAP=50))
    larghi = split_into_chunks(testo)

    assert len(stretti) > len(larghi), "la dimensione configurata non ha effetto"


def test_the_wired_defaults_match_what_was_previously_in_force(monkeypatch):
    """Collegare una manopola non deve cambiare di straforo il comportamento."""
    import config

    assert config.cfg.CHUNK_SIZE == 900
    assert config.cfg.CHUNK_OVERLAP == 140


def test_extracts_html_units_with_headings_and_tables():
    html_content = b"""
    <!DOCTYPE html>
    <html>
    <head><title>Test Page</title><style>.hidden { display: none; }</style></head>
    <body>
        <h1>Regolamento Aziendale</h1>
        <p>Tutti i dipendenti devono seguire le linee guida di sicurezza.</p>
        <h2>Tabella Orari</h2>
        <table>
            <tr><th>Turno</th><th>Orario</th></tr>
            <tr><td>Mattina</td><td>08:00 - 16:30</td></tr>
            <tr><td>Pomeriggio</td><td>14:00 - 22:30</td></tr>
        </table>
    </body>
    </html>
    """
    units = extract_source_units("pagina.html", html_content)
    assert len(units) >= 2
    assert any("linee guida di sicurezza" in u.text for u in units)
    # Verifica che la tabella sia estratta strutturata
    table_unit = next((u for u in units if "Turno | Orario" in u.text), None)
    assert table_unit is not None
    assert "Mattina | 08:00 - 16:30" in table_unit.text
    assert "Tabella 1" in table_unit.locator


def test_extracts_json_and_jsonl_units():
    # 1. JSON strutturato
    json_bytes = b'{"azienda": "Ermes Corp", "dipendenti": 250, "certificazioni": ["ISO27001", "SOC2"]}'
    json_units = extract_source_units("info.json", json_bytes)
    assert len(json_units) >= 3
    assert any("azienda: Ermes Corp" in u.text for u in json_units)
    assert any("certificazioni: ISO27001, SOC2" in u.text for u in json_units)

    # 2. JSONL
    jsonl_bytes = b'{"id": 1, "prodotto": "Server RAG"}\n{"id": 2, "prodotto": "Vector DB"}\n'
    jsonl_units = extract_source_units("catalogo.jsonl", jsonl_bytes)
    assert len(jsonl_units) == 2
    assert "prodotto: Server RAG" in jsonl_units[0].text
    assert jsonl_units[0].locator == "Record JSONL 1"
    assert "prodotto: Vector DB" in jsonl_units[1].text


def test_extracts_docx_paragraphs_and_tables():
    from docx import Document

    doc = Document()
    doc.add_heading("Specifiche Prodotto", level=1)
    doc.add_paragraph("Descrizione generale del componente hardware.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Parametro"
    table.cell(0, 1).text = "Valore"
    table.cell(1, 0).text = "Voltaggio"
    table.cell(1, 1).text = "220V"
    buffer = BytesIO()
    doc.save(buffer)

    units = extract_source_units("specifiche.docx", buffer.getvalue())
    assert len(units) >= 2
    assert any("Descrizione generale del componente hardware." in u.text for u in units)
    table_unit = next((u for u in units if "Parametro | Valore" in u.text), None)
    assert table_unit is not None
    assert "Voltaggio | 220V" in table_unit.text
    assert "Tabella 1" in table_unit.locator
